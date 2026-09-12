"""Merge several DOCX files into one, in the order the user chose.

Nothing here is AI-assisted and nothing is re-typeset: the documents are joined
at the OOXML level, so every paragraph, table, image, footnote and list arrives
in the output exactly as it left its source file.

The work that makes that true is all bookkeeping. Two documents each carry their
own style table, numbering definitions, relationship ids, bookmark names and
drawing ids, and those namespaces collide the moment the XML is concatenated —
identical ids meaning different things, list numbering restarting or running on
where it should not. docxcompose owns that remapping, which is why it is a
dependency rather than a hand-rolled deepcopy loop.

The first document in the order is the base: its styles, page size, margins and
headers/footers govern the merged file, and where a later document defines a
style with the same name but a different body, the base's definition wins. Put
the document whose look should survive first.
"""

import io
import logging
import os

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docxcompose.composer import Composer

logger = logging.getLogger(__name__)

MERGE_SUFFIX = "_merged.docx"

# DOCX is a ZIP container; every one of them starts with the local file header.
ZIP_MAGIC = b"PK\x03\x04"


class DocxMergeError(Exception):
    """The upload cannot be merged: wrong file type, too few files, unreadable."""


def _open(name: str, data: bytes) -> Document:
    """Open one upload as a Document, or explain why it is not one."""
    label = name or "(unnamed file)"

    if not data:
        raise DocxMergeError(f"{label} is empty.")

    # .doc is the old binary format and is not a ZIP — worth its own message,
    # because "not a valid DOCX" sends people looking for a corrupt file.
    if name.lower().endswith(".doc"):
        raise DocxMergeError(
            f"{label} is a legacy .doc file. Save it as .docx in Word first."
        )
    if not name.lower().endswith(".docx"):
        raise DocxMergeError(f"{label} is not a .docx file.")
    if not data.startswith(ZIP_MAGIC):
        raise DocxMergeError(f"{label} is not a valid DOCX file.")

    try:
        return Document(io.BytesIO(data))
    except Exception as exc:
        logger.warning("Could not open %s as a DOCX: %s", label, exc)
        raise DocxMergeError(f"{label} could not be opened as a DOCX.") from exc


def _start_on_new_page(doc: Document, composer: Composer) -> None:
    """Make the next document begin on a fresh page.

    Marking the incoming document's first paragraph `pageBreakBefore` is
    preferable to appending a break paragraph to the base: it adds no empty
    paragraph, so the merged file has no stray blank lines to delete. A document
    that opens with a table has no first paragraph to mark, so that case falls
    back to a break paragraph on the base.
    """
    # The search runs over the body's elements rather than doc.paragraphs because
    # the first entry of doc.paragraphs in a table-led document is the paragraph
    # *after* the table: breaking there would leave the table on the old page.
    for element in doc.element.body:
        if element.tag == qn("w:p"):
            Paragraph(element, doc).paragraph_format.page_break_before = True
            return
        if element.tag == qn("w:tbl"):
            break
        # Anything else up here (sectPr, bookmarkStart, proofState…) is not
        # content, so keep looking.

    composer.doc.add_page_break()


def merge_docx(
    documents: list[tuple[str, bytes]], page_breaks: bool = True
) -> tuple[bytes, str]:
    """Merge `documents` — (filename, bytes) in the order given — into one DOCX.

    With `page_breaks` each document after the first starts on a new page; without
    it, each one continues straight after the previous document's last paragraph.

    Returns the merged file's bytes and a one-line summary for the UI.
    """
    if len(documents) < 2:
        raise DocxMergeError("Choose at least two DOCX files to merge.")

    base_name, base_data = documents[0]
    composer = Composer(_open(base_name, base_data))

    for name, data in documents[1:]:
        doc = _open(name, data)
        if page_breaks:
            _start_on_new_page(doc, composer)
        composer.append(doc)

    buffer = io.BytesIO()
    composer.save(buffer)
    merged = buffer.getvalue()

    summary = f"{len(documents)} documents merged, {_paragraph_count(composer.doc)} paragraphs."
    logger.info("Merged %d documents: %s", len(documents),
                ", ".join(os.path.basename(n) for n, _ in documents))
    return merged, summary


def _paragraph_count(doc: Document) -> int:
    """Non-blank paragraphs in the merged body — a cheap sanity figure for the UI."""
    return sum(1 for p in doc.paragraphs if p.text.strip())
