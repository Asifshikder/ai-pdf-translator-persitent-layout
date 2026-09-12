r"""DOCX merge tests.

Builds its own fixtures with python-docx — no manuals, no API calls — and checks
that merge_docx:
  * keeps every paragraph, and keeps them in the order the caller gave
  * carries images, tables and headings across from the later documents
  * starts each document on a new page when asked, and does not when not
  * refuses a non-DOCX, an empty upload, a .doc and a single file, by message

Run via:
  .\.venv\Scripts\python.exe test_docx_merge.py
"""

import io
import sys

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches

from docx_merge import DocxMergeError, merge_docx

failures = []


def check(name: str, ok: bool, detail: str = "") -> None:
    if ok:
        print(f"  PASS  {name}")
    else:
        print(f"  FAIL  {name}" + (f"\n          {detail}" if detail else ""))
        failures.append(name)


def make_doc(paragraphs: list[str], heading: str | None = None,
             table: bool = False, image: bool = False) -> bytes:
    """A minimal .docx carrying the requested content."""
    doc = Document()
    if heading:
        doc.add_heading(heading, level=1)
    for text in paragraphs:
        doc.add_paragraph(text)
    if table:
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "cell-a"
        t.cell(1, 1).text = "cell-d"
    if image:
        doc.add_picture(io.BytesIO(PNG_1PX), width=Inches(1))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _png_1px() -> bytes:
    """A real 1x1 PNG — python-docx parses the header, so it must be genuine."""
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (1, 1), (220, 40, 40)).save(buf, "PNG")
    return buf.getvalue()


PNG_1PX = _png_1px()


def texts(data: bytes) -> list[str]:
    """Non-blank paragraph texts of a DOCX, in body order."""
    doc = Document(io.BytesIO(data))
    return [p.text for p in doc.paragraphs if p.text.strip()]


def page_break_flags(data: bytes) -> list[bool]:
    """`pageBreakBefore` for every non-blank paragraph, in body order."""
    doc = Document(io.BytesIO(data))
    return [
        p.paragraph_format.page_break_before is True
        for p in doc.paragraphs
        if p.text.strip()
    ]


def image_count(data: bytes) -> int:
    doc = Document(io.BytesIO(data))
    return sum(1 for part in doc.part.package.parts
               if part.partname.startswith("/word/media/"))


print("\nDOCX merge")
print("-" * 60)

# --- order and completeness --------------------------------------------------
a = make_doc(["alpha one", "alpha two"])
b = make_doc(["bravo one"])
c = make_doc(["charlie one", "charlie two", "charlie three"])

merged, summary = merge_docx([("a.docx", a), ("b.docx", b), ("c.docx", c)])
check(
    "every paragraph survives, in the given order",
    texts(merged) == ["alpha one", "alpha two", "bravo one",
                      "charlie one", "charlie two", "charlie three"],
    f"got {texts(merged)}",
)

reordered, _ = merge_docx([("c.docx", c), ("a.docx", a), ("b.docx", b)])
check(
    "reordering the inputs reorders the output",
    texts(reordered) == ["charlie one", "charlie two", "charlie three",
                         "alpha one", "alpha two", "bravo one"],
    f"got {texts(reordered)}",
)

check("summary names the document count", "3 documents merged" in summary, summary)

# --- page breaks -------------------------------------------------------------
flags = page_break_flags(merged)
check(
    "each document after the first starts on a new page",
    flags == [False, False, True, True, False, False],
    f"got {flags}",
)

flat, _ = merge_docx([("a.docx", a), ("b.docx", b)], page_breaks=False)
check(
    "page_breaks=False adds no breaks",
    page_break_flags(flat) == [False, False, False],
    f"got {page_break_flags(flat)}",
)

check(
    "page_breaks=False keeps the same text",
    texts(flat) == ["alpha one", "alpha two", "bravo one"],
    f"got {texts(flat)}",
)

# --- richer content ----------------------------------------------------------
rich = make_doc(["with a table"], heading="Chapter Two", table=True, image=True)
with_rich, _ = merge_docx([("a.docx", a), ("rich.docx", rich)])

check("headings from a later document survive", "Chapter Two" in texts(with_rich),
      f"got {texts(with_rich)}")

doc = Document(io.BytesIO(with_rich))
check("tables from a later document survive", len(doc.tables) == 1,
      f"got {len(doc.tables)} tables")
check("table cell text survives",
      doc.tables and doc.tables[0].cell(0, 0).text == "cell-a")
check("images from a later document survive", image_count(with_rich) == 1,
      f"got {image_count(with_rich)} media parts")

# A document opening with a table has no first paragraph to mark, so the break
# has to fall back to a break paragraph on the base.
table_first = Document()
table_first.add_table(rows=1, cols=1).cell(0, 0).text = "leading table"
table_first.add_paragraph("after the table")
buf = io.BytesIO()
table_first.save(buf)

try:
    table_led, _ = merge_docx([("a.docx", a), ("t.docx", buf.getvalue())])
    body_has_break = any(
        run.find(qn("w:br")) is not None and run.find(qn("w:br")).get(qn("w:type")) == "page"
        for run in Document(io.BytesIO(table_led)).element.body.iter(qn("w:r"))
    )
    check("a document starting with a table still gets a page break", body_has_break)
    check("its content survives", "after the table" in texts(table_led),
          f"got {texts(table_led)}")
except Exception as exc:
    check("a document starting with a table still gets a page break", False, repr(exc))

# --- rejections --------------------------------------------------------------
def rejects(name: str, docs, phrase: str) -> None:
    try:
        merge_docx(docs)
    except DocxMergeError as exc:
        check(name, phrase.lower() in str(exc).lower(), f"message was: {exc}")
    except Exception as exc:
        check(name, False, f"raised {type(exc).__name__}: {exc}")
    else:
        check(name, False, "no error raised")


rejects("a single file is refused", [("a.docx", a)], "at least two")
rejects("an empty list is refused", [], "at least two")
rejects("a PDF is refused", [("a.docx", a), ("b.pdf", b"%PDF-1.7 ...")], "not a .docx")
rejects("a legacy .doc is refused by name",
        [("a.docx", a), ("b.doc", b"\xd0\xcf\x11\xe0")], "legacy")
rejects("an empty upload is refused", [("a.docx", a), ("b.docx", b"")], "empty")
rejects("a .docx that is not a ZIP is refused",
        [("a.docx", a), ("b.docx", b"not a zip at all")], "not a valid docx")

corrupt = bytearray(make_doc(["x"]))
corrupt[200:400] = b"\x00" * 200
rejects("a corrupt .docx is refused", [("a.docx", a), ("b.docx", bytes(corrupt))],
        "could not be opened")

print("-" * 60)
if failures:
    print(f"{len(failures)} failed: " + ", ".join(failures))
    sys.exit(1)
print("all passed")
